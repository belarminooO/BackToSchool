import math
import subprocess
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/belarmino/Documents/BackToSchool")
OUT_DIR = ROOT / "docs"
FIGURES_DIR = OUT_DIR / "figures"
EDITOR_SOURCES_DIR = FIGURES_DIR / "editor_sources"
DOCX_PATH = OUT_DIR / "Relatorio_Back_to_School.docx"
MD_PATH = OUT_DIR / "Relatorio_Back_to_School.md"
TITLE_SANITIZER_SCRIPT = Path("/Users/belarmino/.codex/plugins/cache/openai-primary-runtime/documents/26.601.10930/skills/documents/scripts/google_docs_title_sanitize.py")
ARCHITECTURE_FIGURE = FIGURES_DIR / "arquitetura_sistema.png"
ROUND_FLOW_FIGURE = FIGURES_DIR / "fluxo_ronda.png"
VOICE_PIPELINE_FIGURE = FIGURES_DIR / "pipeline_voz_vigia.png"
LOBBY_FIGURE = FIGURES_DIR / "lobbysample_atual.png"
QUIZ_OVERVIEW_FIGURE = FIGURES_DIR / "quiz_ai_visao_geral.png"
QUIZ_CLASSROOM_FIGURE = FIGURES_DIR / "quiz_ai_classroom2.png"
LOBBY_EDITOR_SOURCE = EDITOR_SOURCES_DIR / "lobbysample_editor.png"
QUIZ_OVERVIEW_SOURCE = EDITOR_SOURCES_DIR / "quiz_ai_visao_geral_editor.png"
QUIZ_CLASSROOM_SOURCE = EDITOR_SOURCES_DIR / "quiz_ai_classroom2_editor.png"
TEMP_LOBBY_EDITOR_SOURCE = FIGURES_DIR / "lobby_sceneview_current.png"
TEMP_QUIZ_OVERVIEW_SOURCE = Path("/var/folders/7l/8jtw8x1j2pgf3mz2xb61xhjm0000gn/T/TemporaryItems/NSIRD_screencaptureui_Imxwmd/Captura de ecrã 2026-06-01, às 20.42.08.png")
TEMP_QUIZ_CLASSROOM_SOURCE = Path("/var/folders/7l/8jtw8x1j2pgf3mz2xb61xhjm0000gn/T/TemporaryItems/NSIRD_screencaptureui_FX0I7D/Captura de ecrã 2026-06-01, às 20.42.25.png")
LOBBY_CROP = (760, 120, 1890, 900)
QUIZ_OVERVIEW_CROP = (430, 120, 1870, 940)
QUIZ_CLASSROOM_CROP = (430, 120, 1870, 940)


COVER_LINES = [
    "Projeto",
    "Ano letivo 2025/2026",
    "ISEL - DEETC - LEIM",
    "",
    "Back to School",
    "1 de junho de 2026",
    "Docente: Prof. Hugo Cordeiro",
    "Arguente: Prof. Leticia Lucas",
    "Trabalho realizado por:",
    "Miguel Ferreira nº: 51878",
    "Belarmino Sacate nº: 52057",
    "Turma: [a confirmar]",
]


REQ_RULES = [
    ("R1.1", 'O sistema deve alternar as rondas entre "Rondas de Perguntas" e "Intervalos nos Corredores da Escola".', "Evidente"),
    ("R1.2", "O sistema deve exibir 5 perguntas interativas por cada ronda no quadro da sala de aula.", "Evidente"),
    ("R1.3", "O sistema deve penalizar o último jogador a chegar à sala de aula ou todos os jogadores que não chegarem à próxima sala dentro do tempo definido.", "Evidente"),
    ("R1.4", "Os jogadores devem ser organizados em equipas de dois, ou sozinhos caso o número de participantes seja ímpar.", "Invisível"),
    ("R1.5", "A equipa ou o jogador com maior pontuação total no final deverá ser o vencedor da partida.", "Evidente"),
    ("R1.6", "O sistema de castigo deve teletransportar o jogador apanhado pelo Vigia para a jaula na sala, impedindo-o de responder a mais perguntas na respetiva ronda.", "Evidente"),
]

REQ_PLAYER = [
    ("R2.1", "Movimento tridimensional em primeira pessoa, bem como salto, sprint e agachamento.", "Evidente"),
    ("R2.2", "Agarrar e empurrar oponentes.", "Evidente"),
    ("R2.3", "Agarrar objetos físicos e atirá-los.", "Evidente"),
    ("R2.4", "Carregar em botões de interação nas secretarias, correspondentes às respostas do quiz, ou interagir com portas.", "Evidente"),
    ("R2.5", "Falar por voz com sistema de proximidade.", "Evidente"),
    ("R2.6", "A câmara com Cinemachine deve seguir os movimentos do jogador em primeira pessoa.", "Evidente"),
]

REQ_PERSISTENCE = [
    ("R3.1", "Guardar os temas do quiz em formato CSV no anfitrião.", "Invisível"),
    ("R3.2", "Gerar e guardar as perguntas do quiz produzidas por inteligência artificial em formato JSON.", "Invisível"),
    ("R3.3", "Guardar as respostas e as respetivas pontuações dos jogadores.", "Invisível"),
]

REQ_INTERFACE = [
    ("R4.1", "Exibir uma mira no centro do ecrã.", "Evidente"),
    ("R4.2", "Mostrar uma barra de tempo que indique visualmente o prazo para responder no quiz e para concluir a corrida nos corredores.", "Evidente"),
    ("R4.3", "Exibir a pontuação individual e coletiva de todos os jogadores através do leaderboard.", "Evidente"),
    ("R4.4", "Mostrar um indicador quando um jogador está a usar o sistema de voz por proximidade.", "Evidente"),
    ("R4.5", "Disponibilizar menus pré-jogo com as opções Create Lobby, Browse Lobby, Join Lobby, Settings e Quit.", "Evidente"),
    ("R4.6", "Disponibilizar definições de criação de lobby para o anfitrião com Round Timer, Round Number e Quiz Themes.", "Evidente"),
]

REQ_AUTONOMY = [
    ("R5.1", "O Vigia deve vigiar a sala de aula com um cone de visão.", "Evidente"),
    ("R5.2", "O Vigia deve deslocar-se com navmesh pela sala durante o quiz.", "Evidente"),
    ("R5.3", "O Vigia deve detetar níveis de som elevados e dirigir-se ao local para aplicar uma penalidade.", "Evidente"),
    ("R5.4", "Os Colegas devem ter um comportamento simples, andando pelos corredores para servirem de obstáculos na corrida.", "Evidente"),
]

REQ_SOUND = [
    ("R6.1", "O jogo deverá reproduzir músicas diferentes consoante os jogadores estejam a responder ao quiz ou a participar na corrida de obstáculos.", "Evidente"),
    ("R6.2", "O sistema deve produzir sons associados a passos, apanhar objetos, cair no chão, atirar objetos, empurrar, agarrar, saltar e pressionar botões.", "Evidente"),
    ("R6.3", "O sistema deve produzir indicadores sonoros associados ao tempo, nomeadamente contagens decrescentes e transições entre rondas.", "Evidente"),
    ("R6.4", "O jogo deve reproduzir som ambiente adequado ao contexto, para além da música.", "Adorno"),
    ("R6.5", "Deve ser possível alterar o volume do som.", "Evidente"),
]

NON_FUNCTIONAL = [
    ("Facilidade de utilização", "Controlos intuitivos para um jogador comum.", "Desejável"),
    ("Facilidade de utilização", "Navegação de menus simples.", "Desejável"),
    ("Acessibilidade", "Complementar ícones com texto.", "Desejável"),
    ("Acessibilidade", "Texto legível.", "Desejável"),
    ("Interação homem-máquina", "Uso exclusivo de rato e teclado.", "Obrigatório"),
    ("Interação homem-máquina", "Resposta rápida do sistema.", "Desejável"),
    ("Plataforma", "Windows.", "Obrigatório"),
    ("Desempenho", "Pelo menos 60 quadros por segundo estáveis.", "Desejável"),
    ("Estética", "Os objetos em jogo devem seguir um mesmo estilo.", "Desejável"),
    ("Estética", "Os menus do jogo devem ser visualmente adequados ao contexto do jogo.", "Desejável"),
    ("Dificuldade do jogo", "O jogo não deve ser propositalmente difícil de completar.", "Desejável"),
]

USE_CASES = [
    ("Mudar opções de jogo", "Permite alterar parâmetros gerais, com destaque para o volume e restantes definições acessíveis através do menu Settings."),
    ("Mudar opções do lobby", "Permite ao anfitrião definir número de rondas, duração das perguntas e temas que servirão de base à geração dos quizzes."),
    ("Responder ao quiz", "Coloca os jogadores numa sala de aula onde devem responder a perguntas antes do tempo terminar, comunicando e sabotando adversários."),
    ("Deslocar até à próxima sala", "Representa a corrida de obstáculos entre rondas, onde a rapidez influencia diretamente a pontuação final."),
    ("Consultar classificação final", "Apresenta o leaderboard atualizado ao longo da partida e expandido no momento de encerramento do jogo."),
]

TECHNOLOGIES = [
    ("Unity 6.0.4f1", "Motor principal do projeto e ambiente de desenvolvimento onde se encontram cenas, prefabs, interface e scripts."),
    ("C#", "Linguagem utilizada para implementar lógica de jogo, sincronização, inteligência artificial e interação."),
    ("PurrNet", "Biblioteca de sincronização multijogador usada para variáveis sincronizadas, chamadas remotas e máquina de estados da partida."),
    ("PurrLobby e Steamworks.NET", "Camada de lobby, descoberta de sessões e integração com o ecossistema Steam."),
    ("MetaVoiceChat", "Solução de voz por proximidade utilizada como base da comunicação entre jogadores."),
    ("Cinemachine", "Gestão das câmaras em primeira pessoa e de câmaras auxiliares de apoio ao desenvolvimento."),
    ("Unity AI Navigation", "Suporte à navegação do Vigia e dos Colegas através de navmesh."),
    ("Input System", "Mapeamento moderno de ações de controlo para movimento, visão, salto e interação."),
    ("Universal Render Pipeline", "Pipeline de renderização adotado pelo projeto para o cenário e para os materiais."),
    ("Gemini 2.5 Flash", "Modelo utilizado para gerar quizzes dinâmicos em formato JSON a partir dos temas definidos pelo anfitrião."),
]

IMPLEMENTATION_STATUS = [
    ("R1.1 / R1.2", "A máquina de estados Waiting, Freeze, Question, ObstacleRace e ChangeRound suporta a alternância de fases e a progressão estruturada das perguntas.", "Implementado"),
    ("R1.4", "O conceito define equipas de dois, mas o ScoreManager atribui atualmente equipas por índice individual, até um máximo de oito equipas distintas.", "Parcial"),
    ("R1.5 / R4.3", "O ScoreManager mantém pontuação sincronizada, atualiza o leaderboard e apresenta um ecrã final de classificação.", "Implementado"),
    ("R2.1", "O movimento em primeira pessoa, a rotação, o salto e o suporte de câmara encontram-se distribuídos por MyCharacterController e pelos controladores associados.", "Implementado"),
    ("R2.2", "O empurrão encontra-se suportado em FirstPersonNetworkState e nos NPCs ColegaEmpurrar; a componente de agarrar não surge com o mesmo grau de consolidação no código principal.", "Parcial"),
    ("R2.3", "A documentação prevê apanhar e atirar objetos, mas a base de código analisada não evidencia uma pipeline tão consolidada como a do empurrão.", "Parcial"),
    ("R2.4", "Interactor, AnswerButton, AnswerButtonManager e ChairInteraction suportam interação contextual com objetos e estações.", "Implementado"),
    ("R2.5 / R4.4 / R5.3", "MetaVoiceChat, NoiseGate, PlayerVoiceState e Vigia articulam voz por proximidade, deteção de ruído e punição do jogador infrator.", "Implementado"),
    ("R3.1 / R3.2", "O anfitrião persiste temas em CSV e perguntas geradas em JSON na área persistente da aplicação.", "Implementado"),
    ("R4.6", "HostQuizManager disponibiliza configuração de rondas, duração e temas do quiz ao anfitrião do lobby.", "Implementado"),
    ("R5.4", "ColegaVaguear e ColegaEmpurrar implementam navegação autónoma e interferência física em contexto de corrida.", "Implementado"),
    ("R6.5", "A infraestrutura de áudio do lobby utiliza PlayerPrefs para preservar níveis de volume entre sessões.", "Implementado"),
]

VALIDATION_SUMMARY = [
    ("Verificação estática do repositório", "Leitura dos scripts principais, dependências, cenas e artefactos de configuração para confirmar o alinhamento entre especificação e protótipo.", "Concluída"),
    ("Console Unity via MCP", "Consulta ao console em 1 de junho de 2026; a amostra devolvida continha apenas registos de arranque do MCP, sem erros recentes nessa observação.", "Concluída com reserva"),
    ("Testes automatizados", "Não foram encontrados diretórios ou suites de testes automatizados dedicados no repositório analisado.", "Em falta"),
    ("Mecanismos defensivos no código", "Existem salvaguardas para ausência de .env, CSV, spawn points, NoiseGate, referências de UI e bindings de cena.", "Parcialmente consolidado"),
    ("Testes multijogador e equilíbrio", "A presença de logs de depuração e de comentários de afinação indica que áreas como empurrão, spawns e integração de cena permanecem em evolução.", "Em progresso"),
]

NEXT_STEPS_PRIORITY = [
    ("Alta", "Consolidar a lógica de equipas de dois e a pontuação coletiva, de forma a alinhar a implementação com o requisito funcional R1.4.", "Elimina um desalinhamento funcional relevante e melhora a justiça competitiva."),
    ("Alta", "Executar ciclos de teste multijogador repetíveis para validar sincronização, transições de ronda, empurrão e penalizações do Vigia.", "Aumenta a confiança no protótipo e reduz riscos de regressão em demonstração."),
    ("Alta", "Reduzir dependências frágeis de cena, revendo referências de inspetor, spawn points, prisões, estações e elementos de interface.", "Diminui falhas por configuração e reforça a robustez da montagem final."),
    ("Média", "Aprofundar as mecânicas de agarrar e de atirar objetos, caso se pretenda manter integralmente a ambição definida no GDD.", "Aproxima o jogo do conceito original e torna a sabotagem mais expressiva."),
    ("Média", "Refinar o polimento audiovisual, a legibilidade do HUD e a coerência estética entre menus, sala de aula e percursos.", "Melhora a perceção de acabamento e a qualidade da apresentação académica."),
]

REFERENCES = [
    "[1] Ferreira, Miguel, e Sacate, Belarmino. Back to School. Apresentação intermédia do projeto, Instituto Superior de Engenharia de Lisboa, maio de 2026.",
    "[2] Ferreira, Miguel, e Sacate, Belarmino. GDD - Back To School. Documento interno de conceção do jogo, 2026.",
    "[3] Ferreira, Miguel, e Sacate, Belarmino. Tradução de GDD para requisitos. Documento interno de análise de requisitos, 2026.",
    "[4] Ferreira, Miguel, e Sacate, Belarmino. Projeto Relatório. Versão preliminar do relatório, 2026.",
    "[5] Código-fonte do projeto Back to School. Repositório local Unity analisado em 1 de junho de 2026.",
    "[6] Documentação técnica das bibliotecas Unity, PurrNet, MetaVoiceChat, Steamworks.NET, Cinemachine e Unity AI Navigation consultada ao longo do desenvolvimento.",
]


CONTENT = [
    {"type": "heading", "level": 1, "title": "1. Introdução"},
    {
        "type": "paragraph",
        "text": "A proposta 61 da Unidade Curricular de Projeto definiu Back to School como um party game multijogador em três dimensões, desenvolvido em Unity para Windows, no qual os jogadores assumem o papel de alunos que competem em questionários temáticos gerados por inteligência artificial. Desde a sua formulação inicial, o projeto distinguiu-se por combinar uma situação escolar reconhecível com uma dinâmica competitiva menos convencional: responder corretamente não basta, porque os jogadores também precisam de comunicar com prudência, evitar a deteção do Vigia e atravessar corredores com obstáculos sempre que a partida transita para uma nova sala.",
    },
    {
        "type": "paragraph",
        "text": "A partir dessa base conceptual, o trabalho evoluiu para um protótipo que articula geração dinâmica de quizzes, sincronização multijogador, interação física, progressão por rondas e voz por proximidade com impacto direto na jogabilidade. O presente relatório documenta essa evolução, relacionando a proposta inicial, o Game Design Document, abreviado GDD, a tradução do GDD para requisitos e o estado efetivamente observado no repositório Unity. A opção por iniciar o desenvolvimento analítico do documento com os requisitos resulta, assim, de uma preocupação de rastreabilidade: antes de discutir soluções técnicas, importa clarificar com precisão aquilo que o sistema se propõe cumprir.",
    },
    {
        "type": "paragraph",
        "text": "O relatório encontra-se organizado da seguinte forma: a Secção 2 sistematiza os requisitos funcionais e não funcionais do projeto; a Secção 3 apresenta a motivação, as inspirações e o conceito final do jogo; a Secção 4 descreve a arquitetura técnica, as tecnologias adotadas e as decisões estruturantes do desenvolvimento; a Secção 5 documenta a implementação atual do protótipo, o grau de conformidade entre requisitos e código, a estratégia de validação seguida e as limitações identificadas; por fim, a Secção 6 sintetiza o trabalho realizado e identifica os passos seguintes considerados mais relevantes. Deste modo, o documento procura acompanhar o percurso do projeto desde a sua formulação conceptual até à sua materialização técnica, sem abdicar de uma leitura crítica do estado atual do sistema.",
    },
    {"type": "heading", "level": 1, "title": "2. Requisitos do sistema"},
    {"type": "heading", "level": 2, "title": "2.1 Enquadramento e objetivos"},
    {
        "type": "paragraph",
        "text": "A formalização de requisitos teve como principal objetivo transformar uma descrição predominantemente lúdica, presente no GDD, numa especificação suficientemente clara para orientar o desenvolvimento, a validação e a redação do presente relatório. Os requisitos foram organizados por áreas funcionais, nomeadamente regras de jogo, controlo do jogador, persistência, interface, autonomia de personagens não jogáveis e som. Foi ainda preservada a categorização usada no documento de origem, distinguindo requisitos evidentes, invisíveis e de adorno, uma vez que essa separação ajuda a perceber se determinada funcionalidade é diretamente observável pelo jogador ou se atua sobretudo ao nível interno do sistema. Em complemento, foram também considerados atributos relacionados com usabilidade, desempenho, acessibilidade, plataforma-alvo e coerência estética.",
    },
    {"type": "heading", "level": 2, "title": "2.2 Requisitos funcionais"},
    {
        "type": "paragraph",
        "text": "Nas tabelas seguintes apresentam-se os requisitos funcionais identificados para o projeto. Em conjunto, estes requisitos definem o ciclo principal da partida, as capacidades do jogador, a infraestrutura mínima de armazenamento, a interface visível ao utilizador, o comportamento autónomo dos agentes presentes no cenário e a componente sonora que suporta a experiência de jogo.",
    },
    {
        "type": "paragraph",
        "text": "Durante esta sistematização foi corrigida uma incoerência de numeração presente no documento intermédio, onde o requisito relativo aos menus pré-jogo surgia identificado como R2.4. No presente relatório, esse requisito passa a ser designado R4.5, por pertencer claramente ao grupo de interface e por essa correção facilitar a leitura do conjunto.",
    },
    {"type": "heading", "level": 3, "title": "2.2.1 Regras de jogo"},
    {"type": "table", "caption": "Requisitos funcionais relativos às regras de jogo.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_RULES, "widths_cm": [2.3, 11.3, 2.4]},
    {"type": "heading", "level": 3, "title": "2.2.2 Controlo do jogador"},
    {"type": "table", "caption": "Requisitos funcionais relativos ao controlo do jogador.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_PLAYER, "widths_cm": [2.3, 11.3, 2.4]},
    {"type": "heading", "level": 3, "title": "2.2.3 Persistência"},
    {"type": "table", "caption": "Requisitos funcionais relativos à persistência de dados.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_PERSISTENCE, "widths_cm": [2.3, 11.3, 2.4]},
    {"type": "heading", "level": 3, "title": "2.2.4 Interface"},
    {"type": "table", "caption": "Requisitos funcionais relativos à interface do utilizador.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_INTERFACE, "widths_cm": [2.3, 11.3, 2.4]},
    {"type": "heading", "level": 3, "title": "2.2.5 Autonomia"},
    {"type": "table", "caption": "Requisitos funcionais relativos à autonomia das personagens não jogáveis.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_AUTONOMY, "widths_cm": [2.3, 11.3, 2.4]},
    {"type": "heading", "level": 3, "title": "2.2.6 Som"},
    {"type": "table", "caption": "Requisitos funcionais relativos ao sistema de som.", "headers": ["Referência", "Descrição", "Categoria"], "rows": REQ_SOUND, "widths_cm": [2.3, 11.3, 2.4]},
    {
        "type": "paragraph",
        "text": "Observa-se que o núcleo funcional do projeto assenta numa alternância rigorosa entre a fase de quiz e a fase de deslocação, sendo essa alternância suportada por mecânicas de resposta, mobilidade, sabotagem e penalização. Para além disso, a presença do sistema de voz por proximidade, do Vigia e dos Colegas transforma o simples questionário num espaço de tensão social e física, o que reforça a identidade própria do jogo. A definição formal destes requisitos foi essencial para alinhar o trabalho de implementação com o conceito inicialmente proposto.",
    },
    {"type": "heading", "level": 2, "title": "2.3 Requisitos não funcionais e atributos do sistema"},
    {
        "type": "paragraph",
        "text": "Os requisitos não funcionais complementam a especificação anterior ao estabelecer restrições e metas de qualidade que condicionam a experiência global do utilizador. Neste projeto, assumem particular importância a facilidade de utilização, a legibilidade dos elementos gráficos, a resposta rápida do sistema, o desempenho estável em Windows e a coerência visual entre menus, objetos e cenários. Ainda que alguns destes atributos sejam classificados como desejáveis e não como obrigatórios, a sua presença influencia diretamente a perceção de polimento, justiça e acessibilidade do jogo.",
    },
    {"type": "table", "caption": "Atributos não funcionais e metas de qualidade do sistema.", "headers": ["Atributo", "Detalhe", "Categoria"], "rows": NON_FUNCTIONAL, "widths_cm": [4.0, 9.5, 2.5]},
    {"type": "heading", "level": 2, "title": "2.4 Casos de utilização principais"},
    {
        "type": "paragraph",
        "text": "Para além da enumeração de requisitos, torna-se útil identificar os casos de utilização que organizam a interação entre jogadores e sistema. A leitura destes casos evidencia os momentos centrais da experiência: configuração da partida, parametrização do lobby, resposta às perguntas, deslocação entre salas e consulta do resultado final. A tabela seguinte resume os principais casos de utilização considerados nesta fase do relatório.",
    },
    {"type": "table", "caption": "Casos de utilização principais considerados no projeto.", "headers": ["Caso de utilização", "Descrição"], "rows": USE_CASES, "widths_cm": [5.0, 11.0]},
    {
        "type": "paragraph",
        "text": "No caso de utilização Responder ao quiz, o sistema inicia uma ronda ao apresentar no quadro a pergunta atual, a contagem decrescente e as opções associadas aos botões físicos disponíveis na secretaria de cada jogador. A partir desse momento, os participantes podem responder, comunicar através de voz por proximidade e perturbar os adversários, tentando simultaneamente evitar a deteção do Vigia. O sistema deve registar a resposta escolhida, indicar quem está a falar e, caso o volume exceda o limiar admissível durante tempo suficiente, desencadear o castigo correspondente. Terminado o tempo da pergunta, a ronda avança para a questão seguinte ou para a fase de transição.",
    },
    {
        "type": "paragraph",
        "text": "No caso de utilização Deslocar até à próxima sala, a porta de saída é aberta e inicia-se um intervalo sob a forma de corrida de obstáculos pelos corredores da escola. Durante esse período, o sistema altera a música, apresenta o tempo restante e permite aos jogadores correr, saltar e interferir uns com os outros enquanto evitam obstáculos e personagens autónomas. Quando todos chegam ao destino ou o tempo termina, a pontuação é revista, sendo penalizados os jogadores que tenham ficado para trás. Este caso de utilização garante a continuidade entre rondas e impede que a experiência fique limitada a uma sucessão estática de perguntas.",
    },
    {"type": "heading", "level": 1, "title": "3. Conceção do jogo"},
    {"type": "heading", "level": 2, "title": "3.1 Motivação e objetivos de design"},
    {
        "type": "paragraph",
        "text": "A motivação principal por detrás de Back to School resultou da vontade de integrar inteligência artificial generativa numa experiência multijogador que não dependesse apenas de reflexos ou apenas de conhecimento. Em vez de construir um jogo de perguntas tradicional, procurou-se combinar resposta cognitiva, improviso social e perturbação física dentro do mesmo ciclo de jogo. Esta combinação revelou-se particularmente interessante porque permite variar continuamente a natureza do desafio sem obrigar o jogador a aprender sistemas demasiado complexos.",
    },
    {
        "type": "paragraph",
        "text": "Do ponto de vista do design, foram definidos quatro objetivos centrais: gerar quizzes dinâmicos de acordo com temas configurados pelo anfitrião, suportar uma experiência multijogador sincronizada e social, misturar a fase de quiz com uma corrida de obstáculos no interior da escola e aumentar a repetibilidade da partida através de conteúdo variável e progressão por rondas. A apresentação intermédia do projeto já evidenciava estes objetivos, colocando a integração de inteligência artificial, o multijogador e a alternância entre sala e corredor no centro da proposta. O conceito acabou, assim, por assentar menos na complexidade individual de cada mecânica e mais na forma como essas mecânicas se encadeiam.",
    },
    {"type": "heading", "level": 2, "title": "3.2 Inspirações"},
    {
        "type": "paragraph",
        "text": "Entre as inspirações mais evidentes encontra-se Kahoot!, cujo modelo de perguntas rápidas, limite temporal e classificação intermédia ajudou a definir a fase de quiz. A ideia de associar cada resposta a uma cor específica e de pressionar os jogadores através do tempo disponível foi transposta para um contexto tridimensional, onde a resposta deixa de ser apenas uma opção num ecrã e passa a corresponder a botões físicos espalhados pela sala. Desta forma, um sistema originalmente pensado para contexto pedagógico serviu de base para uma situação competitiva e caótica.",
    },
    {
        "type": "paragraph",
        "text": "Uma segunda influência relevante foi Dale & Dawson Stationery Supplies, jogo que explora sabotagem, observação social e a presença de uma figura de autoridade. Em Back to School, essa lógica foi adaptada ao ambiente escolar através do Vigia, personagem não jogável que patrulha a sala, investiga ruído e castiga jogadores demasiado ruidosos. A tensão entre colaborar com colegas, enganar adversários e não ser apanhado aproxima o projeto de uma experiência social onde o comportamento dos participantes é tão importante quanto o seu desempenho intelectual.",
    },
    {
        "type": "paragraph",
        "text": "Por fim, Fall Guys contribuiu sobretudo para a fase de deslocação entre salas, marcada por obstáculos, empurrões e competição pelo tempo. A adoção de uma corrida caótica entre rondas permitiu quebrar a cadência do quiz e introduzir uma segunda modalidade de desafio, mais física e imediata. Esta alternância reforça o ritmo da partida e reduz o risco de monotonia, uma vez que cada ronda passa a combinar pressão cognitiva com movimentação no espaço.",
    },
    {"type": "heading", "level": 2, "title": "3.3 Conceito final"},
    {
        "type": "paragraph",
        "text": "O conceito final de Back to School pode ser descrito como um party game competitivo em primeira pessoa, pensado para grupos de amigos e estruturado em rondas compostas por duas fases complementares. Numa primeira fase, os jogadores respondem a perguntas de escolha múltipla numa sala de aula, podendo comunicar, sabotar e ocupar fisicamente o espaço; numa segunda fase, deslocam-se até à sala seguinte através de um percurso com obstáculos, sob pena de perderem pontuação. O enquadramento escolar, aliado à presença do Vigia e de personagens autónomas nos corredores, procura tornar a experiência imediatamente reconhecível, humorística e potencialmente memorável.",
    },
    {
        "type": "paragraph",
        "text": "Esta formulação conceptual introduz um desafio de equilíbrio particularmente relevante: o jogo tem de recompensar simultaneamente conhecimento, coordenação espacial, leitura social do grupo e controlo do risco associado à voz. Se uma destas dimensões dominar excessivamente as restantes, a experiência tende a perder identidade própria, aproximando-se demasiado de um quiz tradicional ou, pelo contrário, de uma simples corrida caótica. Por essa razão, a análise do protótipo não deve limitar-se à existência de funcionalidades isoladas, devendo considerar de que modo essas funcionalidades se articulam para sustentar a proposta lúdica inicialmente definida.",
    },
    {"type": "heading", "level": 1, "title": "4. Arquitetura e tecnologias"},
    {"type": "heading", "level": 2, "title": "4.1 Visão geral da arquitetura"},
    {
        "type": "paragraph",
        "text": "A implementação do projeto assenta sobre Unity 6, utilizando uma arquitetura de rede peer-to-peer, abreviada P2P, suportada por PurrNet e pelos serviços de lobby integrados via Steam. Nesta arquitetura, o anfitrião da sessão assume responsabilidades adicionais, nomeadamente a leitura dos temas do quiz, a geração das perguntas com recurso a uma interface de programação de aplicações externa e a disseminação do estado necessário aos restantes pares. Esta opção permite manter a lógica da partida sincronizada sem recorrer a uma infraestrutura dedicada de servidor central.",
    },
    {
        "type": "paragraph",
        "text": "Ao nível lógico, o sistema encontra-se organizado em subsistemas relativamente autónomos: gestão do estado da partida, geração de quizzes, pontuação e progressão, controlo do jogador, comunicação por voz, inteligência artificial do Vigia e dos Colegas, interface e configuração de lobby. A separação destes subsistemas foi importante para reduzir acoplamento entre funcionalidades, permitir iteração incremental e tornar o código mais legível. Embora se trate ainda de um protótipo em evolução, a estrutura do projeto já revela uma tentativa consistente de dividir responsabilidades.",
    },
    {"type": "heading", "level": 2, "title": "4.2 Tecnologias principais"},
    {
        "type": "paragraph",
        "text": "O conjunto de tecnologias identificado no repositório confirma a orientação apresentada anteriormente na demonstração intermédia do projeto. Para além do motor e da linguagem base, foram integradas bibliotecas específicas para sincronização multijogador, voz, navegação de agentes, câmaras e gestão de input. Essa composição tecnológica não foi arbitrária; cada componente responde diretamente a uma necessidade concreta do desenho do jogo.",
    },
    {"type": "table", "caption": "Tecnologias principais adotadas no desenvolvimento de Back to School.", "headers": ["Tecnologia", "Finalidade no projeto"], "rows": TECHNOLOGIES, "widths_cm": [4.2, 11.8]},
    {
        "type": "paragraph",
        "text": "A presença simultânea destas ferramentas evidencia uma abordagem orientada para a integração de serviços e bibliotecas especializadas, em vez da reimplementação de funcionalidades base. Tal opção acelera o desenvolvimento do protótipo e concentra o esforço do grupo na lógica própria do jogo. Em contrapartida, exige maior cuidado na compatibilização entre pacotes e na gestão do estado de rede, particularmente quando a voz e a geração de conteúdo dinâmico passam a influenciar diretamente a jogabilidade.",
    },
    {"type": "heading", "level": 2, "title": "4.3 Estrutura lógica do projeto"},
    {
        "type": "paragraph",
        "text": "O repositório encontra-se organizado de forma coerente com essa divisão. A pasta Assets/Quiz concentra a lógica específica das rondas, dos botões, da interface do quadro, da pontuação e da inteligência artificial do contexto escolar; a pasta Assets/_Scripts agrega scripts transversais relacionados com jogador, voz, câmaras e comportamento em rede; a pasta Assets/Settings centraliza a configuração do anfitrião; por sua vez, o pacote de lobby e as bibliotecas de terceiros permanecem separadas do código principal. Esta estrutura facilita a manutenção do projeto e clarifica quais os componentes que constituem desenvolvimento próprio e quais os que correspondem a dependências externas.",
    },
    {"type": "heading", "level": 2, "title": "4.4 Decisões técnicas estruturantes"},
    {
        "type": "paragraph",
        "text": "Entre as decisões técnicas mais relevantes destaca-se a concentração de responsabilidades sensíveis no anfitrião da partida. Essa opção observa-se na geração do quiz, na leitura dos temas configurados, no armazenamento local dos ficheiros gerados e na difusão do estado necessário aos restantes participantes. A solução foi particularmente adequada porque reduz duplicação de chamadas externas, evita divergência no conteúdo gerado e simplifica a sincronização do início de partida, ainda que aumente a dependência do host para operações críticas.",
    },
    {
        "type": "paragraph",
        "text": "Uma segunda decisão importante foi a separação entre transporte de áudio e interpretação do comportamento vocal. Em vez de associar diretamente a mecânica do Vigia ao próprio fluxo bruto de voz, o sistema introduz uma camada intermédia de filtragem e deteção, composta pelo NoiseGate e pelo PlayerVoiceState. Esta decomposição melhora a clareza arquitetural, porque o envio de áudio continua a ser tratado como problema de comunicação em rede, enquanto a deteção de infrações passa a ser tratada como problema de lógica de jogo. Do mesmo modo, o uso combinado de SyncVars e chamadas remotas permite distinguir entre estado persistente, como pontuações e nome do estado atual, e eventos de atualização imediata, como sincronização da pergunta apresentada ou teletransporte de jogadores.",
    },
    {
        "type": "paragraph",
        "text": "A articulação global destes subsistemas pode ser observada na Figura 1, onde se resume a forma como o anfitrião, a máquina de estados, a camada de rede, a geração de conteúdo e os subsistemas de jogador comunicam entre si. Embora o diagrama simplifique inevitavelmente algumas dependências de implementação, ele ajuda a perceber que o projeto não assenta num único script centralizador, mas antes num conjunto de componentes especializados que cooperam para manter a partida consistente.",
    },
    {"type": "figure", "path": ARCHITECTURE_FIGURE, "caption": "Arquitetura lógica simplificada do sistema Back to School.", "widths_cm": 16.2},
    {"type": "heading", "level": 2, "title": "4.5 Fluxo técnico de uma ronda"},
    {
        "type": "paragraph",
        "text": "Para além da visão estática da arquitetura, importa compreender o fluxo operacional de uma ronda completa. A sessão começa numa fase de espera, durante a qual o anfitrião valida a presença dos jogadores e a disponibilidade do quiz. De seguida, o sistema entra num pequeno estado de preparação e avança para a sequência de perguntas, onde cada iteração sincroniza quadro, temporizador, resposta correta e feedback para todos os clientes. Quando a última pergunta da ronda termina, a lógica transita para a corrida de obstáculos, atualiza o registo de chegadas e calcula penalizações antes de decidir se a partida continua ou se deve ser apresentado o resultado final.",
    },
    {
        "type": "paragraph",
        "text": "Este encadeamento pode ser observado na Figura 2. A sua relevância para o relatório reside no facto de demonstrar que a dinâmica do jogo foi pensada como um ciclo repetível e verificável, e não como uma sucessão ad hoc de eventos. Em termos de engenharia, esta opção favorece manutenção, depuração e expansão futura, uma vez que cada transição de estado passa a corresponder a uma responsabilidade explícita e mais fácil de testar.",
    },
    {"type": "figure", "path": ROUND_FLOW_FIGURE, "caption": "Fluxo simplificado de uma ronda de jogo e respetivas transições de estado.", "widths_cm": 16.2},
    {"type": "heading", "level": 1, "title": "5. Implementação atual do protótipo"},
    {"type": "heading", "level": 2, "title": "5.1 Geração dinâmica de quizzes"},
    {
        "type": "paragraph",
        "text": "A geração dinâmica de quizzes encontra-se concentrada no componente GenerateQuizJSON, o qual lê os temas previamente definidos pelo anfitrião a partir do ficheiro quiz_info.csv guardado na área persistente da aplicação. Com base nesses temas e no número de rondas configurado, o sistema constrói programaticamente um pedido textual que impõe restrições explícitas ao modelo generativo, tais como o número exato de rondas, a existência de cinco perguntas por ronda, a presença de quatro opções por pergunta e a randomização do índice da resposta correta. Esta preparação do pedido revelou-se fundamental para reduzir ambiguidade e aumentar a previsibilidade da resposta.",
    },
    {
        "type": "paragraph",
        "text": "Depois de enviado o pedido à API do modelo Gemini 2.5 Flash, a resposta recebida é limpa, validada e convertida para uma estrutura interna de dados antes de ser guardada em formato JavaScript Object Notation, abreviado JSON, no armazenamento local do anfitrião. A mesma componente atualiza ainda a pergunta corrente, as opções de resposta e o índice correto, tornando essas informações imediatamente disponíveis para a interface do quiz e para a avaliação de pontuação. O código inclui também lógica de repetição da chamada em caso de falha de parsing e prevê um modo alternativo baseado em quiz local quando a geração dinâmica é ignorada, o que simplifica a depuração e reduz dependência de chamadas externas durante testes internos.",
    },
    {"type": "heading", "level": 2, "title": "5.2 Gestão do ciclo de jogo"},
    {
        "type": "paragraph",
        "text": "O ciclo principal da partida é orquestrado pelo QuizGameManager, que funciona como ponto de coordenação entre pontuação, geração de quiz, pontos de spawn, temporização e interface. Em vez de concentrar toda a lógica num único fluxo monolítico, o projeto recorre a uma máquina de estados da biblioteca PurrNet, composta pelos estados Waiting, Freeze, Question, ObstacleRace e ChangeRound. Cada estado possui responsabilidade delimitada: aguardar jogadores e quiz, apresentar uma pausa curta antes da pergunta, gerir a fase de resposta, controlar a corrida de obstáculos e preparar a ronda seguinte.",
    },
    {
        "type": "paragraph",
        "text": "Esta solução permite que a progressão da partida seja previsível, sincronizada e facilmente extensível. O estado inicial só arranca a sessão quando o quiz já foi gerado e quando o número esperado de jogadores se encontra registado e instanciado na cena. A transição para a fase de pergunta sincroniza o conteúdo do quadro com todos os clientes, enquanto a passagem para a fase de corrida redefine o temporizador, limpa o registo de chegadas e reposiciona os jogadores fora da sala. Por fim, o estado de mudança de ronda atualiza a ronda corrente, volta a atribuir estações e prepara a nova pergunta, fechando o ciclo.",
    },
    {"type": "heading", "level": 2, "title": "5.3 Jogador, interação e mobilidade"},
    {
        "type": "paragraph",
        "text": "Do ponto de vista do jogador, a base de controlo foi implementada sobre Kinematic Character Controller, através do script MyCharacterController, que assegura movimento em primeira pessoa, rotação desacoplada da câmara, salto e gestão de sensibilidade do rato. A integração com o novo Input System de Unity permite separar claramente as ações de movimento, olhar, salto e alternância de câmara, o que facilita futura reconfiguração. Para além disso, foram incluídos mecanismos de teletransporte e de feedback visual, nomeadamente uma vinheta que informa o jogador sobre a correção ou o erro da resposta submetida.",
    },
    {
        "type": "paragraph",
        "text": "A interação com o espaço é suportada pelo componente Interactor, que utiliza raycasts a partir da fonte de interação do jogador para detetar objetos interativos, apresentar indicação contextual no ecrã e executar a ação adequada quando a tecla de interação é pressionada. Os botões de resposta foram modelados através da dupla AnswerButton e AnswerButtonManager, permitindo associar uma estação a um jogador específico, registar a opção escolhida e animar a pressão do botão em todos os clientes. Já as interações físicas de perturbação assentam em componentes de rede dedicados, com destaque para FirstPersonNetworkState, responsável por coordenar a mecânica de empurrão entre pares e por aplicar knockback de forma sincronizada. Embora a documentação conceptual mencione igualmente agarrar e atirar objetos, o código principal analisado evidencia o empurrão como a interação competitiva mais consolidada nesta fase do desenvolvimento.",
    },
    {"type": "heading", "level": 2, "title": "5.4 Comunicação por voz e deteção do Vigia"},
    {
        "type": "paragraph",
        "text": "Uma das componentes mais distintivas do projeto é a utilização de voz por proximidade como mecânica central e não apenas como acessório social. A integração entre MetaVoiceChat e PurrNet é realizada pelo componente PurrNetNetProvider, responsável por encaminhar quadros de áudio através de chamadas remotas não fiáveis, adequadas à natureza contínua e efémera deste tipo de comunicação. No lado do jogador, o PlayerVoiceController gere o silenciamento do microfone e os indicadores visuais associados ao estado da transmissão.",
    },
    {
        "type": "paragraph",
        "text": "Para que a voz possa influenciar a jogabilidade, o sistema inclui ainda um filtro de ruído próprio, designado NoiseGate, que calcula o volume aproximado em decibéis, descarta amostras abaixo do limiar definido e disponibiliza um valor suavizado para análise. O componente PlayerVoiceState transforma essa leitura num estado sincronizado de fala, o qual pode ser observado pelo servidor e pelos restantes subsistemas. Esta etapa intermédia foi importante, uma vez que desacopla o transporte de áudio da lógica de deteção comportamental.",
    },
    {
        "type": "paragraph",
        "text": "A partir dessa informação, o Vigia implementa um comportamento reativo baseado em patrulha por waypoints, cone de visão, raio de audição, investigação de origem sonora e progressão de deteção até à captura. Quando a fala é suficientemente intensa e persistente, o agente desloca-se até ao local suspeito e, se confirmar o infrator em condições válidas, envia-o para a jaula da sala, retirando-lhe a possibilidade de continuar a responder nessa ronda. O resultado é uma mecânica interessante de risco e recompensa, onde comunicar demais pode beneficiar a equipa no imediato mas prejudicá-la pouco depois.",
    },
    {
        "type": "paragraph",
        "text": "A sequência de transformação do áudio em sinal jogável encontra-se sintetizada na Figura 3. Esta representação é útil porque evidencia que o sistema não trata a voz apenas como canal de comunicação, mas como fonte de dados sujeita a filtragem, sincronização e interpretação comportamental. Tal escolha constitui um dos aspetos tecnicamente mais distintivos do projeto, distinguindo-o de soluções multijogador onde a voz é apenas um serviço paralelo sem impacto mecânico direto.",
    },
    {"type": "figure", "path": VOICE_PIPELINE_FIGURE, "caption": "Pipeline simplificada da voz por proximidade até à deteção do Vigia.", "widths_cm": 16.0},
    {"type": "heading", "level": 2, "title": "5.5 Pontuação, teletransporte e progressão"},
    {
        "type": "paragraph",
        "text": "A gestão de pontuação e progressão entre espaços é assegurada sobretudo pelos componentes ScoreManager, SpawnManager e RoomTrigger. O primeiro mantém valores sincronizados de pontuação por equipa, regista os jogadores presentes, avalia respostas corretas, aplica penalizações a quem não chega atempadamente à sala seguinte e atualiza o leaderboard tanto durante a partida como no ecrã final. O segundo define pontos de spawn por ronda para o interior da sala, para o exterior e para a prisão, executando teletransportes sincronizados sempre que a fase de jogo assim o exige.",
    },
    {
        "type": "paragraph",
        "text": "Durante a corrida de obstáculos, os jogadores que entram no gatilho da sala correta são assinalados como chegados, enquanto os restantes permanecem elegíveis para penalização quando o temporizador termina. Para além disso, o sistema aplica cosméticos simples de destaque aos melhores e piores classificados, reforçando a leitura competitiva do estado da partida. Apesar de a lógica de equipas ainda poder ser refinada para alinhar totalmente com a regra conceptual de duplas, a infraestrutura existente já suporta classificação, progressão e encerramento da sessão de forma coerente.",
    },
    {"type": "heading", "level": 2, "title": "5.6 Interface e menus"},
    {
        "type": "paragraph",
        "text": "Ao nível de interface, o projeto combina elementos específicos do quiz com a infraestrutura de menus fornecida pelo sistema de lobby. O componente HostQuizManager disponibiliza ao anfitrião controlos para número de rondas, duração das perguntas e temas do quiz, persistindo essas escolhas localmente antes do início da partida. Por sua vez, a interface do quiz é atualizada pelo QuizUIManager, que ativa o quadro correspondente à ronda atual, apresenta pergunta e opções e ajusta a visibilidade do cursor consoante o estado da partida. O temporizador visual é gerido por RoundTimerManager, que adapta tanto o valor apresentado como a cor de preenchimento ao estado corrente da partida.",
    },
    {
        "type": "paragraph",
        "text": "A solução de menus suporta criação de lobby, procura de salas existentes, entrada por código e acesso a definições, estando em linha com o que havia sido previsto no GDD e apresentado na demonstração intermédia. Em conjunto com indicadores de voz, leaderboard, temporizador e feedback visual de resposta, esta camada de interface procura manter o jogador informado sem sobrecarregar o ecrã. O grau de acabamento visual ainda pode evoluir, mas a estrutura funcional dos menus e do heads-up display, abreviado HUD, encontra-se montada.",
    },
    {
        "type": "paragraph",
        "text": "A materialização atual desta infraestrutura de arranque pode ser observada na Figura 4, correspondente a uma vista recente da cena LobbySample no editor. Ainda que se trate de uma vista de trabalho e não de uma captura promocional, ela é particularmente útil para o relatório, porque documenta o estado real dos elementos de menu, dos painéis de configuração e da organização espacial usada no lobby multijogador.",
    },
    {"type": "figure", "path": LOBBY_FIGURE, "caption": "Vista atual da cena LobbySample, usada para menus e configuração inicial do lobby.", "widths_cm": 16.0},
    {"type": "heading", "level": 2, "title": "5.7 Estado atual do protótipo"},
    {
        "type": "paragraph",
        "text": "No estado atual do protótipo, já se observam implementações concretas para a geração de quizzes por inteligência artificial, a sincronização multijogador, a progressão por rondas, a fase de sala de aula, a fase de corrida, a deteção do Vigia e a apresentação do resultado final. A apresentação intermédia já apontava como conquistas principais a integração de inteligência artificial no jogo e a estabilidade multijogador via Steam e PurrNet, e a leitura do repositório confirma que esses dois eixos continuam a constituir o núcleo do projeto. Em contrapartida, permanecem relevantes tarefas de polimento, equilíbrio de parâmetros, consolidação de interações físicas, melhoria visual dos cenários, correção de erros remanescentes e aprofundamento de testes em contexto real de partida.",
    },
    {
        "type": "paragraph",
        "text": "Do ponto de vista espacial, a cena Quiz AI já evidencia uma estrutura global concreta, com várias salas, percursos de ligação, obstáculos e elementos de ambientação, como se pode observar na Figura 5. Esta vista geral é útil porque mostra que o projeto já não se encontra circunscrito a uma única sala isolada, existindo uma organização de cenário compatível com a alternância entre perguntas e deslocação competitiva prevista nos requisitos.",
    },
    {"type": "figure", "path": QUIZ_OVERVIEW_FIGURE, "caption": "Vista geral atual da cena Quiz AI, evidenciando a disposição global do percurso e das salas.", "widths_cm": 16.0},
    {
        "type": "paragraph",
        "text": "Por sua vez, a Figura 6 destaca uma vista recente da Classroom 2, onde já se reconhecem o quadro de apresentação, as secretarias distribuídas pela sala e os botões físicos usados para responder ao quiz. Em conjunto, estas duas figuras documentam com maior fidelidade o estado atual do protótipo do que as capturas antigas anteriormente usadas, refletindo melhor a configuração que se encontra hoje disponível no editor.",
    },
    {"type": "figure", "path": QUIZ_CLASSROOM_FIGURE, "caption": "Vista atual da Classroom 2 na cena Quiz AI, com quadro interativo, secretarias e botões de resposta.", "widths_cm": 16.0},
    {"type": "heading", "level": 2, "title": "5.8 Conformidade entre requisitos e implementação"},
    {
        "type": "paragraph",
        "text": "Uma vez que se trata de um projeto de fim de curso, não basta descrever funcionalidades isoladas; importa também avaliar o grau em que a implementação observada satisfaz os requisitos inicialmente definidos. A tabela seguinte resume esse cruzamento para os aspetos mais relevantes, distinguindo requisitos já suportados de forma clara, requisitos parcialmente implementados e áreas cuja presença no conceito ainda não corresponde a uma solução plenamente estabilizada no código analisado.",
    },
    {"type": "table", "caption": "Matriz resumida de conformidade entre requisitos e implementação observada.", "headers": ["Referência", "Evidência observada", "Estado"], "rows": IMPLEMENTATION_STATUS, "widths_cm": [3.2, 10.4, 2.4]},
    {
        "type": "paragraph",
        "text": "A leitura desta matriz permite retirar uma conclusão importante: o protótipo já cumpre com solidez a espinha dorsal do jogo, nomeadamente a alternância entre fases, a geração de quizzes, a sincronização da partida, a gestão de voz e a presença do Vigia. No entanto, também revela desalinhamentos que devem ser assumidos de forma transparente, como a regra de equipas de dois ainda não refletida integralmente na lógica atual de ScoreManager e a menor maturidade de mecânicas como agarrar ou atirar objetos quando comparadas com o sistema de empurrão. Esta honestidade analítica valoriza o relatório, porque demonstra compreensão técnica do estado real do projeto.",
    },
    {"type": "heading", "level": 2, "title": "5.9 Validação e estratégia de teste"},
    {
        "type": "paragraph",
        "text": "A validação do protótipo, nesta fase, assenta sobretudo em verificação manual, depuração em contexto de execução e análise direta do comportamento dos subsistemas principais. A própria organização do código evidencia esta realidade: vários componentes incluem mensagens de erro, mensagens de aviso e caminhos de fallback destinados a detetar rapidamente referências em falta, ficheiros inexistentes, dependências de cena não configuradas ou incoerências de spawn. Esta abordagem é compatível com um protótipo em evolução, embora ainda não substitua uma estratégia formal e repetível de testes automatizados.",
    },
    {
        "type": "paragraph",
        "text": "A consulta do console Unity via MCP, efetuada em 1 de junho de 2026, não revelou erros recentes na amostra observada, contendo apenas registos de inicialização do próprio serviço MCP. Ainda assim, essa observação não deve ser interpretada como prova exaustiva de estabilidade, mas apenas como um indicador pontual de que não havia falhas imediatas evidentes no momento da inspeção. Em paralelo, a pesquisa realizada no repositório não revelou suites dedicadas de testes automatizados, o que significa que a confiança no sistema depende sobretudo de ensaios manuais e da robustez dos mecanismos de verificação embutidos nos próprios scripts.",
    },
    {"type": "table", "caption": "Síntese das evidências de validação identificadas no projeto.", "headers": ["Área de validação", "Observação", "Estado"], "rows": VALIDATION_SUMMARY, "widths_cm": [4.0, 9.6, 2.4]},
    {"type": "heading", "level": 2, "title": "5.10 Limitações técnicas e trabalho em aberto"},
    {
        "type": "paragraph",
        "text": "A análise do projeto torna igualmente visíveis algumas limitações que importa registar. Em primeiro lugar, várias funcionalidades dependem de configuração correta de cena e de referências no inspetor, como spawn points, estações, pontos de prisão, UI boards e componentes de voz. A presença de mensagens explícitas para falta dessas referências mostra que o grupo teve consciência do problema e tentou reduzir o impacto em tempo de execução, mas também indica que o sistema ainda está sensível a erros de montagem do cenário. Para um protótipo académico esta situação é compreensível, embora deva ser progressivamente reduzida até à entrega final.",
    },
    {
        "type": "paragraph",
        "text": "Em segundo lugar, subsistem áreas onde a intenção de design é mais ambiciosa do que a maturidade atual da implementação. A mecânica de equipas deveria organizar jogadores em duplas, mas a lógica de atribuição atual não impõe esse comportamento da forma prevista no documento de requisitos. Do mesmo modo, a documentação conceptual sugere um leque de interações físicas mais vasto do que aquele que se observa claramente consolidado no código principal, onde o empurrão surge como funcionalidade mais desenvolvida. Estas diferenças não anulam o mérito do protótipo, mas exigem priorização cuidadosa das próximas iterações.",
    },
    {
        "type": "paragraph",
        "text": "Por fim, a ausência de uma bateria de testes automatizados e a necessidade continuada de afinação em contexto multijogador sugerem que o principal desafio das etapas seguintes não será apenas acrescentar funcionalidades, mas estabilizar e validar aquilo que já existe. Nesse sentido, o trabalho em aberto passa por consolidar a lógica de equipas, reforçar testes repetíveis, reduzir dependências frágeis de cena, melhorar a coerência entre requisitos e implementação e refinar a experiência audiovisual e de interação. Este conjunto de ações tende a ter impacto direto tanto na qualidade do jogo como na solidez da apresentação académica do projeto.",
    },
    {"type": "heading", "level": 2, "title": "5.11 Priorização do trabalho até à entrega final"},
    {
        "type": "paragraph",
        "text": "Numa perspetiva de entrega final, não basta reconhecer limitações; importa ordenar o esforço de acordo com o impacto académico e funcional de cada intervenção. A priorização apresentada na tabela seguinte resulta da análise cruzada entre requisitos, estado atual do código e risco técnico observado. O objetivo não é apenas aumentar o número de funcionalidades, mas sobretudo reforçar a coerência entre aquilo que o projeto promete, aquilo que já demonstra e aquilo que será defendido no relatório e na apresentação final.",
    },
    {"type": "table", "caption": "Prioridades de evolução identificadas para a fase final do projeto.", "headers": ["Prioridade", "Ação", "Impacto esperado"], "rows": NEXT_STEPS_PRIORITY, "widths_cm": [2.3, 8.7, 5.0]},
    {
        "type": "paragraph",
        "text": "A principal leitura desta priorização é que a etapa seguinte do projeto deve privilegiar consolidação sobre expansão indiscriminada. Em particular, a correção dos desalinhamentos mais visíveis, a estabilização da experiência multijogador e a redução de fragilidade de cena tendem a produzir ganhos mais relevantes do que a introdução apressada de novas mecânicas. Esta perspetiva é especialmente importante num projeto de fim de curso, uma vez que a avaliação valoriza não apenas ambição conceptual, mas também consistência técnica, capacidade de análise e maturidade de execução.",
    },
    {"type": "heading", "level": 1, "title": "6. Conclusão"},
    {
        "type": "paragraph",
        "text": "Back to School evoluiu de uma ideia centrada em quizzes para um protótipo multijogador mais ambicioso, no qual conhecimento, movimentação, sabotagem e comunicação coexistem no mesmo ciclo de jogo. A introdução formal dos requisitos no início do relatório permitiu clarificar o âmbito funcional do projeto e criar uma ponte mais sólida entre o material conceptual do GDD e o desenvolvimento concreto observado no repositório. Esta clarificação revelou-se importante para enquadrar o trabalho já realizado e para identificar, com maior objetividade, os aspetos que ainda necessitam de consolidação.",
    },
    {
        "type": "paragraph",
        "text": "Em termos técnicos, o projeto já demonstra uma base consistente: existe uma arquitetura de rede funcional, um mecanismo de geração automática de conteúdo, um ciclo de jogo estruturado por estados e uma integração expressiva da voz na própria jogabilidade. Simultaneamente, a análise crítica do código mostra que a qualidade de um relatório final não depende apenas da enumeração do que já funciona, mas também da capacidade de reconhecer desalinhamentos, fragilidades e prioridades de evolução. O trabalho futuro deverá concentrar-se no reforço da robustez, no refinamento da experiência de utilizador, na revisão de equilíbrio entre mecânicas, na consolidação da lógica de equipas e na formalização de estratégias de teste mais sistemáticas. Ainda assim, o estado atual do Back to School mostra que o conceito definido no início do projeto não só é exequível, como já se encontra materializado num protótipo jogável, tecnicamente coerente e suficientemente promissor para justificar investimento adicional de polimento e validação.",
    },
    {"type": "heading", "level": 1, "title": "7. Referências"},
    *({"type": "paragraph", "text": ref} for ref in REFERENCES),
]


SECTION_OVERVIEWS = {
    "2. Requisitos do sistema": "A presente secção organiza a especificação do sistema em quatro planos complementares. Começa por enquadrar a origem e os objetivos dos requisitos, prossegue com a identificação dos requisitos funcionais, reúne depois os requisitos não funcionais e termina com os casos de utilização principais que sintetizam a interação entre jogadores e sistema.",
    "3. Conceção do jogo": "Nesta secção apresenta-se a base conceptual do projeto antes da análise técnica detalhada. As subseções seguintes expõem a motivação e os objetivos de design, discutem as principais inspirações que moldaram a proposta e explicitam o conceito final que orienta o protótipo atualmente desenvolvido.",
    "4. Arquitetura e tecnologias": "Esta secção descreve a estrutura técnica que suporta o protótipo e clarifica a forma como os principais subsistemas se articulam. Nas subseções seguintes apresentam-se a visão global da arquitetura, as tecnologias adotadas, a organização lógica do projeto, as decisões técnicas estruturantes e o fluxo operacional de uma ronda completa.",
    "5. Implementação atual do protótipo": "A presente secção incide sobre a materialização concreta do projeto no repositório e no editor Unity. Para esse efeito, as subseções seguintes analisam a geração dinâmica de quizzes, a gestão do ciclo de jogo, o jogador e a interação, a voz e o Vigia, a pontuação e progressão, a interface, o estado atual do protótipo, a conformidade com os requisitos, a validação, as limitações técnicas e as prioridades até à entrega final.",
}


HEADING_PAGE_MAP = {
    "1. Introdução": 5,
    "2. Requisitos do sistema": 5,
    "2.1 Enquadramento e objetivos": 5,
    "2.2 Requisitos funcionais": 6,
    "2.2.1 Regras de jogo": 6,
    "2.2.2 Controlo do jogador": 7,
    "2.2.3 Persistência": 7,
    "2.2.4 Interface": 8,
    "2.2.5 Autonomia": 8,
    "2.2.6 Som": 9,
    "2.3 Requisitos não funcionais e atributos do sistema": 9,
    "2.4 Casos de utilização principais": 10,
    "3. Conceção do jogo": 11,
    "3.1 Motivação e objetivos de design": 11,
    "3.2 Inspirações": 12,
    "3.3 Conceito final": 12,
    "4. Arquitetura e tecnologias": 13,
    "4.1 Visão geral da arquitetura": 13,
    "4.2 Tecnologias principais": 13,
    "4.3 Estrutura lógica do projeto": 14,
    "4.4 Decisões técnicas estruturantes": 15,
    "4.5 Fluxo técnico de uma ronda": 16,
    "5. Implementação atual do protótipo": 16,
    "5.1 Geração dinâmica de quizzes": 17,
    "5.2 Gestão do ciclo de jogo": 17,
    "5.3 Jogador, interação e mobilidade": 18,
    "5.4 Comunicação por voz e deteção do Vigia": 18,
    "5.5 Pontuação, teletransporte e progressão": 19,
    "5.6 Interface e menus": 20,
    "5.7 Estado atual do protótipo": 21,
    "5.8 Conformidade entre requisitos e implementação": 23,
    "5.9 Validação e estratégia de teste": 24,
    "5.10 Limitações técnicas e trabalho em aberto": 25,
    "5.11 Priorização do trabalho até à entrega final": 26,
    "6. Conclusão": 27,
    "7. Referências": 28,
}

FIGURE_PAGE_MAP = {
    "Arquitetura lógica simplificada do sistema Back to School.": 15,
    "Fluxo simplificado de uma ronda de jogo e respetivas transições de estado.": 16,
    "Pipeline simplificada da voz por proximidade até à deteção do Vigia.": 19,
    "Vista atual da cena LobbySample, usada para menus e configuração inicial do lobby.": 21,
    "Vista geral atual da cena Quiz AI, evidenciando a disposição global do percurso e das salas.": 22,
    "Vista atual da Classroom 2 na cena Quiz AI, com quadro interativo, secretarias e botões de resposta.": 22,
}

TABLE_PAGE_MAP = {
    "Requisitos funcionais relativos às regras de jogo.": 6,
    "Requisitos funcionais relativos ao controlo do jogador.": 7,
    "Requisitos funcionais relativos à persistência de dados.": 7,
    "Requisitos funcionais relativos à interface do utilizador.": 8,
    "Requisitos funcionais relativos à autonomia das personagens não jogáveis.": 8,
    "Requisitos funcionais relativos ao sistema de som.": 9,
    "Atributos não funcionais e metas de qualidade do sistema.": 10,
    "Casos de utilização principais considerados no projeto.": 10,
    "Tecnologias principais adotadas no desenvolvimento de Back to School.": 13,
    "Matriz resumida de conformidade entre requisitos e implementação observada.": 23,
    "Síntese das evidências de validação identificadas no projeto.": 25,
    "Prioridades de evolução identificadas para a fase final do projeto.": 26,
}


def ensure_output_dir() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    EDITOR_SOURCES_DIR.mkdir(parents=True, exist_ok=True)


def get_font(size: int, bold: bool = False):
    candidates = []
    if bold:
        candidates.extend(
            [
                "DejaVuSans-Bold.ttf",
                "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                "/Library/Fonts/Arial Bold.ttf",
            ]
        )
    else:
        candidates.extend(
            [
                "DejaVuSans.ttf",
                "/System/Library/Fonts/Supplemental/Arial.ttf",
                "/Library/Fonts/Arial.ttf",
            ]
        )

    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def fit_font_to_width(draw: ImageDraw.ImageDraw, text: str, max_width: int, *, preferred_size: int, minimum_size: int, bold: bool = False):
    for size in range(preferred_size, minimum_size - 1, -1):
        font = get_font(size, bold=bold)
        text_width = draw.textbbox((0, 0), text, font=font)[2]
        if text_width <= max_width:
            return font
    return get_font(minimum_size, bold=bold)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]

    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_box(
    draw: ImageDraw.ImageDraw,
    *,
    x: int,
    y: int,
    width: int,
    title: str,
    body: str,
    fill: str,
    outline: str,
    accent: str,
    title_size: int = 24,
    body_size: int = 18,
) -> tuple[int, int, int, int]:
    radius = 24
    padding_x = 18
    padding_y = 16
    title_font = fit_font_to_width(draw, title, width - (padding_x * 2), preferred_size=title_size, minimum_size=18, bold=True)
    body_font = get_font(body_size, bold=False)
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    available_width = width - (padding_x * 2)
    lines = wrap_text(draw, body, body_font, available_width)
    line_height = draw.textbbox((0, 0), "Ag", font=body_font)[3] + 5
    body_height = max(line_height * len(lines), line_height)
    header_height = max(52, (title_bbox[3] - title_bbox[1]) + 18)
    total_height = header_height + padding_y + body_height + padding_y
    x1, y1, x2, y2 = x, y, x + width, y + total_height

    draw.rounded_rectangle((x1, y1, x2, y2), radius=radius, fill=fill, outline=outline, width=3)
    draw.rounded_rectangle((x1, y1, x2, y1 + header_height + 10), radius=radius, fill=accent, outline=accent)
    title_y = y1 + (header_height - (title_bbox[3] - title_bbox[1])) / 2 - 2
    draw.text((x1 + padding_x, title_y), title, fill="white", font=title_font)

    start_y = y1 + header_height + padding_y - 2
    for line in lines:
        draw.text((x1 + padding_x, start_y), line, fill="#1F2937", font=body_font)
        start_y += line_height

    return x1, y1, x2, y2


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], *, label: str | None = None, color: str = "#0F172A") -> None:
    width = 6
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    arrow_size = 16
    p1 = (
        end[0] - arrow_size * math.cos(angle - math.pi / 6),
        end[1] - arrow_size * math.sin(angle - math.pi / 6),
    )
    p2 = (
        end[0] - arrow_size * math.cos(angle + math.pi / 6),
        end[1] - arrow_size * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, p1, p2], fill=color)
    if label:
        font = get_font(18, bold=False)
        bbox = draw.textbbox((0, 0), label, font=font)
        mid_x = (start[0] + end[0] - (bbox[2] - bbox[0])) / 2
        mid_y = (start[1] + end[1] - (bbox[3] - bbox[1])) / 2 - 24
        draw.rounded_rectangle((mid_x - 8, mid_y - 4, mid_x + bbox[2] - bbox[0] + 8, mid_y + bbox[3] - bbox[1] + 4), radius=8, fill="white")
        draw.text((mid_x, mid_y), label, fill=color, font=font)


def add_canvas_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    title_font = get_font(36, bold=True)
    subtitle_font = get_font(20, bold=False)
    draw.text((70, 48), title, fill="#0F172A", font=title_font)
    draw.text((70, 98), subtitle, fill="#475569", font=subtitle_font)
    draw.line([(70, 138), (1530, 138)], fill="#CBD5E1", width=3)


def generate_architecture_figure(path: Path) -> None:
    image = Image.new("RGB", (1600, 920), "white")
    draw = ImageDraw.Draw(image)
    add_canvas_title(
        draw,
        "Arquitetura lógica do Back to School",
        "Relação simplificada entre anfitrião, estado de jogo, geração de quizzes e subsistemas de jogabilidade.",
    )

    lobby_box = draw_box(draw, x=80, y=200, width=340, title="Lobby e configuração", body="PurrLobby, Steamworks.NET e definições do anfitrião para rondas, tempo e temas.", fill="#ECFDF5", outline="#10B981", accent="#10B981")
    host_box = draw_box(draw, x=80, y=500, width=340, title="Serviços do anfitrião", body="Leitura de CSV, chamada ao Gemini, escrita de JSON e gestão de dados persistentes.", fill="#FEF3C7", outline="#F59E0B", accent="#F59E0B")
    core_box = draw_box(draw, x=540, y=300, width=480, title="Coordenação da partida", body="QuizGameManager e máquina de estados PurrNet asseguram Waiting, Freeze, Question, ObstacleRace e ChangeRound.", fill="#EFF6FF", outline="#2563EB", accent="#2563EB", title_size=23)
    player_box = draw_box(draw, x=1170, y=170, width=350, title="Jogador e interação", body="Movimento, raycasts, botões de resposta, teletransporte e feedback local.", fill="#F8FAFC", outline="#64748B", accent="#64748B")
    voice_box = draw_box(draw, x=1170, y=395, width=350, title="Voz e deteção", body="MetaVoiceChat, NoiseGate, PlayerVoiceState e lógica reativa do Vigia.", fill="#FDF2F8", outline="#DB2777", accent="#DB2777")
    score_box = draw_box(draw, x=1170, y=640, width=350, title="Pontuação e progressão", body="ScoreManager, SpawnManager, RoomTrigger e leaderboard sincronizado.", fill="#EEF2FF", outline="#7C3AED", accent="#7C3AED")

    draw_arrow(draw, (lobby_box[2], (lobby_box[1] + lobby_box[3]) // 2), (core_box[0], core_box[1] + 70), label="parâmetros")
    draw_arrow(draw, (host_box[2], (host_box[1] + host_box[3]) // 2), (core_box[0], core_box[3] - 70), label="quiz e dados")
    draw_arrow(draw, (core_box[2], core_box[1] + 75), (player_box[0], player_box[1] + 55), label="estado")
    draw_arrow(draw, (core_box[2], (core_box[1] + core_box[3]) // 2), (voice_box[0], voice_box[1] + 78), label="eventos")
    draw_arrow(draw, (core_box[2], core_box[3] - 60), (score_box[0], score_box[1] + 75), label="pontuação")
    draw_arrow(draw, ((voice_box[0] + voice_box[2]) // 2, voice_box[3]), ((score_box[0] + score_box[2]) // 2, score_box[1]), label="penalizações")

    image.save(path)


def generate_round_flow_figure(path: Path) -> None:
    image = Image.new("RGB", (1600, 930), "white")
    draw = ImageDraw.Draw(image)
    add_canvas_title(
        draw,
        "Fluxo técnico de uma ronda",
        "Sequência resumida de estados que organiza a sessão desde a preparação inicial até à decisão de continuar ou terminar.",
    )

    waiting_box = draw_box(draw, x=70, y=250, width=230, title="Waiting", body="Jogadores e quiz prontos.", fill="#ECFDF5", outline="#10B981", accent="#10B981")
    freeze_box = draw_box(draw, x=330, y=250, width=240, title="Freeze", body="Pausa curta e preparação da pergunta.", fill="#FEF3C7", outline="#F59E0B", accent="#F59E0B")
    question_box = draw_box(draw, x=600, y=250, width=270, title="Question", body="Cinco perguntas com quadro, botões e temporizador.", fill="#EFF6FF", outline="#2563EB", accent="#2563EB")
    race_box = draw_box(draw, x=900, y=250, width=290, title="ObstacleRace", body="Corrida entre salas e registo de chegadas.", fill="#FDF2F8", outline="#DB2777", accent="#DB2777", title_size=22)
    change_box = draw_box(draw, x=1220, y=250, width=250, title="ChangeRound", body="Atualiza ronda, estações e spawns.", fill="#EEF2FF", outline="#7C3AED", accent="#7C3AED", title_size=22)

    for start_box, end_box, label in [
        (waiting_box, freeze_box, "pronto"),
        (freeze_box, question_box, "inicia ronda"),
        (question_box, race_box, "última pergunta"),
        (race_box, change_box, "avaliar resultados"),
    ]:
        draw_arrow(draw, (start_box[2], (start_box[1] + start_box[3]) // 2), (end_box[0], (end_box[1] + end_box[3]) // 2), label=label)

    decision_box = draw_box(draw, x=560, y=545, width=550, title="Decisão de continuidade", body="Caso existam rondas por jogar, a partida regressa ao estado Freeze; caso contrário, apresenta a classificação final.", fill="#F8FAFC", outline="#0F172A", accent="#0F172A", title_size=22)
    leaderboard_box = draw_box(draw, x=1290, y=545, width=250, title="Leaderboard final", body="Ordenação das equipas e encerramento da sessão.", fill="#ECFEFF", outline="#0891B2", accent="#0891B2", title_size=21)

    draw_arrow(draw, ((change_box[0] + change_box[2]) // 2, change_box[3]), (decision_box[2] - 40, decision_box[1]), label="fim da ronda")
    draw_arrow(draw, (decision_box[0] + 20, decision_box[1] + 70), (freeze_box[0] + 100, freeze_box[3]), label="há mais rondas")
    draw_arrow(draw, (decision_box[2], (decision_box[1] + decision_box[3]) // 2), (leaderboard_box[0], (leaderboard_box[1] + leaderboard_box[3]) // 2), label="fim da partida")

    image.save(path)


def generate_voice_pipeline_figure(path: Path) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    add_canvas_title(
        draw,
        "Voz por proximidade e deteção do Vigia",
        "Transformação do áudio do jogador em estado sincronizado e, quando aplicável, em penalização jogável.",
    )

    mic_box = draw_box(draw, x=55, y=290, width=220, title="Microfone", body="Captação local do áudio do jogador.", fill="#ECFEFF", outline="#0891B2", accent="#0891B2")
    meta_box = draw_box(draw, x=310, y=290, width=240, title="MetaVoiceChat", body="Gestão do fluxo de voz e estado de transmissão.", fill="#EEF2FF", outline="#4F46E5", accent="#4F46E5", title_size=22)
    provider_box = draw_box(draw, x=585, y=290, width=260, title="PurrNetNetProvider", body="Envio do áudio por chamadas remotas adequadas a dados efémeros.", fill="#EFF6FF", outline="#2563EB", accent="#2563EB", title_size=20)
    gate_box = draw_box(draw, x=880, y=290, width=230, title="NoiseGate", body="Filtra ruído e estima intensidade vocal.", fill="#FEF3C7", outline="#D97706", accent="#D97706")
    state_box = draw_box(draw, x=1145, y=290, width=250, title="PlayerVoiceState", body="Sincroniza o estado de fala relevante para a lógica.", fill="#FDF2F8", outline="#DB2777", accent="#DB2777", title_size=20)
    guard_box = draw_box(draw, x=1185, y=585, width=310, title="Vigia e punição", body="Investiga, confirma infração e envia o jogador para a jaula.", fill="#FEF2F2", outline="#DC2626", accent="#DC2626", title_size=22)

    for left, right, label in [
        (mic_box, meta_box, "transmissão"),
        (meta_box, provider_box, "encaminhamento"),
        (provider_box, gate_box, "amostras"),
        (gate_box, state_box, "volume"),
    ]:
        draw_arrow(draw, (left[2], (left[1] + left[3]) // 2), (right[0], (right[1] + right[3]) // 2), label=label)

    draw_arrow(draw, ((state_box[0] + state_box[2]) // 2, state_box[3]), ((guard_box[0] + guard_box[2]) // 2, guard_box[1]), label="suspeita")
    image.save(path)


def update_editor_source(temp_path: Path, stable_path: Path) -> None:
    if stable_path.exists():
        return
    if temp_path.exists():
        shutil.copyfile(temp_path, stable_path)
    else:
        raise FileNotFoundError(f"Não foi possível localizar a captura de origem em '{temp_path}' nem em '{stable_path}'.")


def crop_editor_scene(source_path: Path, output_path: Path, crop_box: tuple[int, int, int, int]) -> None:
    image = Image.open(source_path).convert("RGB")
    cropped = image.crop(crop_box)
    cropped.save(output_path)


def prepare_figures() -> None:
    ensure_output_dir()
    generate_architecture_figure(ARCHITECTURE_FIGURE)
    generate_round_flow_figure(ROUND_FLOW_FIGURE)
    generate_voice_pipeline_figure(VOICE_PIPELINE_FIGURE)
    update_editor_source(TEMP_LOBBY_EDITOR_SOURCE, LOBBY_EDITOR_SOURCE)
    update_editor_source(TEMP_QUIZ_OVERVIEW_SOURCE, QUIZ_OVERVIEW_SOURCE)
    update_editor_source(TEMP_QUIZ_CLASSROOM_SOURCE, QUIZ_CLASSROOM_SOURCE)
    crop_editor_scene(LOBBY_EDITOR_SOURCE, LOBBY_FIGURE, LOBBY_CROP)
    crop_editor_scene(QUIZ_OVERVIEW_SOURCE, QUIZ_OVERVIEW_FIGURE, QUIZ_OVERVIEW_CROP)
    crop_editor_scene(QUIZ_CLASSROOM_SOURCE, QUIZ_CLASSROOM_FIGURE, QUIZ_CLASSROOM_CROP)


def set_rtl_safe_font(run, name: str, size_pt: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:cs"), name)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    width_twips = int(Cm(width_cm).twips)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Cm(width_cm)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def enable_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def configure_page(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def add_page_number_footer(section) -> None:
    section.different_first_page_header_footer = True
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_rtl_safe_font(run, "Cambria", 10)
    add_field(paragraph, "PAGE", "1")


def create_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.3

    for style_name, size, before, after in [
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 12, 10, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Cambria"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("1F1F1F")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        style._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "CoverLine" not in doc.styles:
        cover_style = doc.styles.add_style("CoverLine", WD_STYLE_TYPE.PARAGRAPH)
        cover_style.base_style = normal
        cover_style.font.name = "Cambria"
        cover_style.font.size = Pt(12)
        cover_style.font.bold = False
        cover_style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        cover_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        cover_style._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
        cover_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover_style.paragraph_format.space_after = Pt(4)

    if "ReferenceLine" not in doc.styles:
        ref_style = doc.styles.add_style("ReferenceLine", WD_STYLE_TYPE.PARAGRAPH)
        ref_style.base_style = normal
        ref_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ref_style.paragraph_format.space_after = Pt(6)

    if "SectionLead" not in doc.styles:
        lead_style = doc.styles.add_style("SectionLead", WD_STYLE_TYPE.PARAGRAPH)
        lead_style.base_style = normal
        lead_style.font.name = "Cambria"
        lead_style.font.size = Pt(11)
        lead_style.font.italic = False
        lead_style.font.bold = False
        lead_style.font.color.rgb = RGBColor.from_string("000000")
        lead_style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        lead_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        lead_style._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
        lead_style.paragraph_format.space_before = Pt(0)
        lead_style.paragraph_format.space_after = Pt(8)
        lead_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, left_indent, is_bold in [("IndexLevel1", 0.0, True), ("IndexLevel2", 0.7, False), ("IndexLevel3", 1.4, False)]:
        if style_name in doc.styles:
            continue
        index_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        index_style.base_style = normal
        index_style.font.name = "Cambria"
        index_style.font.size = Pt(11)
        index_style.font.bold = is_bold
        index_style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
        index_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
        index_style._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
        index_style.paragraph_format.space_before = Pt(0)
        index_style.paragraph_format.space_after = Pt(1)
        index_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        index_style.paragraph_format.left_indent = Cm(left_indent)
        index_style.paragraph_format.first_line_indent = Cm(0)
        tabs = index_style.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    caption = doc.styles["Caption"]
    caption.font.name = "Cambria"
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    caption._element.rPr.rFonts.set(qn("w:cs"), "Cambria")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    spacer_added = False
    for line in COVER_LINES:
        paragraph = doc.add_paragraph(style="CoverLine")
        if not spacer_added:
            paragraph.paragraph_format.space_before = Pt(12)
        if line == "":
            paragraph.paragraph_format.space_after = Pt(18)
            spacer_added = True
            continue
        run = paragraph.add_run(line)
        size = 12
        bold = False
        if line == "Back to School":
            size = 18
            bold = True
            paragraph.paragraph_format.space_before = Pt(18)
            paragraph.paragraph_format.space_after = Pt(10)
        elif line in {"Projeto", "Ano letivo 2025/2026", "ISEL - DEETC - LEIM"}:
            bold = True if line == "Projeto" else False
        elif line == "Trabalho realizado por:":
            bold = True
            paragraph.paragraph_format.space_before = Pt(8)
        set_rtl_safe_font(run, "Cambria", size, bold=bold)


def add_heading(doc: Document, title: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(title)
    set_rtl_safe_font(run, "Cambria", {1: 16, 2: 13, 3: 12}[level], bold=True)


def add_paragraph(doc: Document, text: str, style_name: str = "Normal") -> None:
    paragraph = doc.add_paragraph(style=style_name)
    run = paragraph.add_run(text)
    size = 11
    set_rtl_safe_font(run, "Cambria", size)


def add_section_overview(doc: Document, title: str) -> None:
    overview = SECTION_OVERVIEWS.get(title)
    if overview:
        add_paragraph(doc, overview, style_name="SectionLead")


def add_caption(doc: Document, label: str, text: str, number: int | None = None) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    prefix = f"{label} {number}: " if number is not None else f"{label}: "
    run = paragraph.add_run(f"{prefix}{text}")
    set_rtl_safe_font(run, "Cambria", 10, bold=False)


def add_figure(doc: Document, path: Path, caption: str, width_cm: float, number: int) -> None:
    add_caption(doc, "Figura", caption, number=number)
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")


def add_table(doc: Document, headers, rows, widths_cm) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    header_cells = table.rows[0].cells
    for idx, (cell, header, width) in enumerate(zip(header_cells, headers, widths_cm)):
        set_cell_width(cell, width)
        set_cell_margins(cell)
        shade_cell(cell, "EDEFF2")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(header))
        set_rtl_safe_font(run, "Cambria", 10.5, bold=True)

    for row in rows:
        row_cells = table.add_row().cells
        for idx, (cell, value, width) in enumerate(zip(row_cells, row, widths_cm)):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_rtl_safe_font(run, "Cambria", 10.5)

    doc.add_paragraph("")


def get_table_captions() -> list[str]:
    return [item["caption"] for item in CONTENT if item["type"] == "table"]


def get_figure_captions() -> list[str]:
    return [item["caption"] for item in CONTENT if item["type"] == "figure"]


def add_index_entry(doc: Document, text: str, page_number: int, *, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"IndexLevel{level}")
    run = paragraph.add_run(f"{text}\t{page_number}")
    set_rtl_safe_font(run, "Cambria", 11, bold=(level == 1))


def add_static_index_entries(doc: Document, label: str, captions: list[str], page_map: dict[str, int]) -> None:
    if not captions:
        add_paragraph(doc, f"Não existem {label.lower()} legendadas nesta versão do relatório.")
        return

    for idx, caption in enumerate(captions, start=1):
        add_index_entry(doc, f"{label} {idx}: {caption}", page_map[caption], level=1)


def get_heading_index_entries() -> list[tuple[str, int]]:
    return [(item["title"], item["level"]) for item in CONTENT if item["type"] == "heading"]


def add_front_matter_indices(doc: Document) -> None:
    add_heading(doc, "Índice", 1)
    for title, level in get_heading_index_entries():
        add_index_entry(doc, title, HEADING_PAGE_MAP[title], level=level)
    doc.add_page_break()

    add_heading(doc, "Índice de figuras", 1)
    add_static_index_entries(doc, "Figura", get_figure_captions(), FIGURE_PAGE_MAP)
    doc.add_page_break()

    add_heading(doc, "Índice de tabelas", 1)
    add_static_index_entries(doc, "Tabela", get_table_captions(), TABLE_PAGE_MAP)
    doc.add_page_break()


def build_markdown() -> str:
    lines = []
    lines.extend(COVER_LINES)
    lines.append("")
    lines.append("## Índice")
    lines.append("")
    for title, _level in get_heading_index_entries():
        lines.append(f"{title} .... {HEADING_PAGE_MAP[title]}")
    lines.append("")
    lines.append("## Índice de figuras")
    lines.append("")
    if get_figure_captions():
        for idx, caption in enumerate(get_figure_captions(), start=1):
            lines.append(f"Figura {idx}: {caption} .... {FIGURE_PAGE_MAP[caption]}")
    else:
        lines.append("Não aplicável nesta versão do relatório, uma vez que o corpo do documento ainda não inclui figuras legendadas.")
    lines.append("")
    lines.append("## Índice de tabelas")
    lines.append("")
    for idx, caption in enumerate(get_table_captions(), start=1):
        lines.append(f"Tabela {idx}: {caption} .... {TABLE_PAGE_MAP[caption]}")
    lines.append("")

    table_counter = 0
    figure_counter = 0
    for item in CONTENT:
        if item["type"] == "heading":
            lines.append(f'{"#" * item["level"]} {item["title"]}')
            lines.append("")
            overview = SECTION_OVERVIEWS.get(item["title"])
            if overview:
                lines.append(overview)
                lines.append("")
        elif item["type"] == "paragraph":
            lines.append(item["text"])
            lines.append("")
        elif item["type"] == "figure":
            figure_counter += 1
            figure_path = Path(item["path"]).resolve()
            lines.append(f"Figura {figure_counter}: {item['caption']}")
            lines.append("")
            lines.append(f"![Figura {figure_counter}]({figure_path})")
            lines.append("")
        elif item["type"] == "table":
            table_counter += 1
            lines.append(f"Tabela {table_counter}: {item['caption']}")
            lines.append("")
            headers = item["headers"]
            rows = item["rows"]
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for row in rows:
                lines.append("| " + " | ".join(str(col) for col in row) + " |")
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def build_docx() -> Document:
    doc = Document()
    enable_update_fields(doc)
    configure_page(doc.sections[0])
    add_page_number_footer(doc.sections[0])
    create_styles(doc)

    core = doc.core_properties
    core.title = "Relatório Back to School"
    core.author = "Miguel Ferreira; Belarmino Sacate"
    core.subject = "Projeto Back to School"
    core.comments = "Versão de trabalho gerada automaticamente a partir do estado atual do projeto."

    add_cover(doc)
    doc.add_page_break()
    add_front_matter_indices(doc)

    figure_counter = 0
    table_counter = 0
    for item in CONTENT:
        if item["type"] == "heading":
            add_heading(doc, item["title"], item["level"])
            add_section_overview(doc, item["title"])
        elif item["type"] == "paragraph":
            style = "ReferenceLine" if item["text"].startswith("[") else "Normal"
            add_paragraph(doc, item["text"], style_name=style)
        elif item["type"] == "figure":
            figure_counter += 1
            add_figure(doc, Path(item["path"]), item["caption"], item["widths_cm"], figure_counter)
        elif item["type"] == "table":
            table_counter += 1
            add_caption(doc, "Tabela", item["caption"], number=table_counter)
            add_table(doc, item["headers"], item["rows"], item["widths_cm"])

    return doc


def sanitize_google_docs_docx(input_docx: Path, output_docx: Path) -> None:
    subprocess.run(
        [sys.executable, str(TITLE_SANITIZER_SCRIPT), str(input_docx), "--out", str(output_docx)],
        check=True,
    )
    subprocess.run(
        [sys.executable, str(TITLE_SANITIZER_SCRIPT), str(output_docx), "--check"],
        check=True,
    )


def main() -> None:
    prepare_figures()
    MD_PATH.write_text(build_markdown(), encoding="utf-8")
    doc = build_docx()
    with tempfile.NamedTemporaryFile(suffix=".docx", prefix="back_to_school_raw_", delete=False) as handle:
        raw_docx = Path(handle.name)
    try:
        doc.save(raw_docx)
        sanitize_google_docs_docx(raw_docx, DOCX_PATH)
    finally:
        if raw_docx.exists():
            raw_docx.unlink()
    print(f"Markdown: {MD_PATH}")
    print(f"DOCX: {DOCX_PATH}")


if __name__ == "__main__":
    main()
